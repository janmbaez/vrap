import csv
import io
import zipfile
import json
from datetime import date, datetime
from openpyxl import load_workbook
from ..schemas import FindingCreate

MAX_BYTES = 100 * 1024 * 1024
MAX_ROWS = 100000
FIELDS = ['hostname', 'ip', 'name', 'plugin_id', 'cves', 'cvss', 'vpr', 'severity', 'port', 'protocol', 'asset_tags', 'asset_criticality', 'business_criticality', 'data_classification', 'regulatory', 'exposure', 'production', 'environment', 'existing_controls', 'exploit_available', 'kev', 'first_seen', 'last_seen']

def parse_file(filename, mime, content):
    if len(content) > MAX_BYTES:
        raise ValueError('File exceeds 100 MiB')
    if filename.lower().endswith('.csv'):
        if mime not in ('text/csv', 'application/csv', 'application/vnd.ms-excel', 'text/plain'):
            raise ValueError('CSV MIME type is not supported')
        text = content.decode('utf-8-sig')
        if '\x00' in text:
            raise ValueError('Binary content is not CSV')
        rows = list(csv.reader(io.StringIO(text)))
    elif filename.lower().endswith('.xlsx'):
        if mime != 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
            raise ValueError('XLSX MIME type is not supported')
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(x.file_size for x in entries) > 300 * 1024 * 1024:
                raise ValueError('Expanded workbook is too large')
            if any('vbaproject' in x.filename.lower() for x in entries):
                raise ValueError('Macro-enabled content is not allowed')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.active
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 100:
            raise ValueError('Maximum 100,000 rows and 100 columns')
        rows = []
        for row in sheet.iter_rows():
            if any(c.data_type == 'f' for c in row):
                raise ValueError('Formulas are not allowed; upload values only')
            rows.append([c.value.isoformat() if isinstance(c.value, (datetime, date)) else (str(c.value) if c.value is not None else '') for c in row])
        workbook.close()
    else:
        raise ValueError('Only CSV and XLSX are supported')
    if not rows or not rows[0] or len(rows) > MAX_ROWS + 1 or max(map(len, rows)) > 100:
        raise ValueError('Empty file or exceeds 100,000 rows / 100 columns')
    headers = [str(h).strip() for h in rows[0]]
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('Column headers must be nonempty and unique')
    result = []
    for i, row in enumerate(rows[1:], 2):
        if not any(str(c).strip() for c in row):
            continue
        if len(row) != len(headers):
            raise ValueError(f'Row {i} has a different number of columns')
        if any(len(str(c)) > 30000 for c in row):
            raise ValueError(f'Row {i} contains an oversized cell')
        result.append((i, dict(zip(headers, row))))
    return headers, result

def boolean(value):
    normalized = value.strip().lower().replace('-', '_').replace(' ', '_')
    if normalized in ('true', 'yes', 'y', '1', 'available', 'exploits_are_available'):
        return True
    if normalized in ('false', 'no', 'n', '0', 'not_available', 'unavailable', 'no_known_exploits_are_available'):
        return False
    if normalized in ('not_required', 'no_exploit_is_required', 'unknown', 'not_applicable', 'n/a', 'na'):
        return None
    raise ValueError('Boolean values must be Yes/No, True/False, 1/0, or a supported Tenable availability value')

def date_only(value):
    candidate = value.strip()
    try:
        return datetime.fromisoformat(candidate.replace('Z', '+00:00')).date().isoformat()
    except ValueError:
        try:
            return date.fromisoformat(candidate).isoformat()
        except ValueError:
            raise ValueError(f'Invalid ISO date or timestamp: {candidate}')

TAG_CATEGORIES = {
    'location': 'location', 'device type': 'device_type', 'environment': 'environment',
    'application': 'application', 'patching group': 'patching_group', 'cis baseline': 'compliance_baseline',
    'audit': 'audit_scope', 'exception': 'exception_state', 'change owner': 'change_owner',
    'support owner': 'support_owner', 'zengrc issue': 'zengrc_issue', 'websites': 'website_scope',
    'issue': 'issue_state', 'device exceptions': 'device_exception', 'os': 'os'
}

def asset_tags(value):
    """Convert Tenable JSON tags (or simple delimited text) to stable tags."""
    raw = value.strip()
    if not raw: return []
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        parsed = [part.strip() for part in raw.replace(';', ',').split(',') if part.strip()]
    result = []
    for item in parsed if isinstance(parsed, list) else []:
        if isinstance(item, dict):
            category = TAG_CATEGORIES.get(str(item.get('category', '')).strip().casefold(), str(item.get('category', '')).strip().casefold().replace(' ', '_'))
            value = str(item.get('value', '')).strip()
            if category and value: result.append(f'{category}:{value}')
        elif str(item).strip():
            result.append(str(item).strip())
    return list(dict.fromkeys(result))[:100]

def normalize(original, mapping):
    values = {field: str(original.get(column, '')).strip() for field, column in mapping.items() if column and str(original.get(column, '')).strip()}
    context = {}
    for key in ('asset_criticality', 'business_criticality', 'data_classification', 'exposure', 'environment'):
        if key in values:
            context[key] = values.pop(key)
    if 'regulatory' in values:
        context['regulatory'] = [v.strip() for v in values.pop('regulatory').replace(';', ',').split(',')]
    if 'production' in values:
        context['production'] = boolean(values.pop('production'))
        context.setdefault('environment', 'Production' if context['production'] else 'Development')
    if 'existing_controls' in values:
        context['other_controls'] = values.pop('existing_controls') + ' (Imported presence; not validated)'
    if 'asset_tags' in values:
        values['asset_tags'] = asset_tags(values['asset_tags'])
    for key in ('exploit_available', 'kev'):
        if key in values:
            parsed = boolean(values[key])
            if parsed is None:
                values.pop(key)
            else:
                values[key] = parsed
    for key in ('first_seen', 'last_seen'):
        if key in values:
            values[key] = date_only(values[key])
    if 'protocol' in values:
        values['protocol'] = values['protocol'].lower()
    if 'cves' in values:
        values['cves'] = [v.strip() for v in values['cves'].replace(';', ',').split(',')]
    values['hostname'] = values.get('hostname') or values.get('ip', '')
    return FindingCreate(**values, context=context)
